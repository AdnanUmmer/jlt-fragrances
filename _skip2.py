from openpyxl import load_workbook
ws=load_workbook(r"..\..\JLT_Fragrances_List.xlsx").active
hdr=[c.value for c in ws[1]]
print(ws.max_row)
for r in range(813,822):
    vals=[c.value for c in ws[r]]
    print(r, [ (h,v) for h,v in zip(hdr,vals) if v not in (None,"")])
rows={}
for r in range(2,813):
    d=dict(zip(hdr,[c.value for c in ws[r]]))
    k=(str(d["Name"]).strip().lower(),str(d["Brand Inspiration"]).strip().lower())
    rows.setdefault(k,[]).append((r,d))
for k,v in rows.items():
    if len(v)>1:
        a,b=v[0][1],v[1][1]
        diff=[h for h in hdr if h!="Sl No" and a[h]!=b[h]]
        print(k,[x[0] for x in v],"differs in:",diff)

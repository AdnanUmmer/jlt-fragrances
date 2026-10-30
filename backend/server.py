from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")
# Development overrides never load in production.
import os
if os.environ.get("APP_ENV", "development") != "production":
    load_dotenv(ROOT_DIR / ".env.local")

import os
import io
import csv
import json
import uuid
import logging
import hmac
import hashlib
import smtplib
import re
from email.message import EmailMessage
from datetime import datetime, timezone, timedelta
from typing import List, Optional

import bcrypt
import jwt
import razorpay
from requests.exceptions import ConnectionError as PaymentConnectionError, Timeout as PaymentTimeout
from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, UploadFile, File, Query
from fastapi.responses import PlainTextResponse, Response
from xml.sax.saxutils import escape as xml_escape
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from openpyxl import load_workbook

from seed_helpers import build_product, slugify, SIZES
from discovery import normalized, matches, filter_options, key
from quiz import inventory, matching, choices, STEPS
from offers import COMBOS, DISCOVERY_SETS
import copy
from starlette.concurrency import run_in_threadpool

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

def payment_configuration_error():
    key_id = os.environ.get("RAZORPAY_KEY_ID", "").strip()
    secret = os.environ.get("RAZORPAY_KEY_SECRET", "").strip()
    if not key_id or not secret:
        return "missing RAZORPAY_KEY_ID or RAZORPAY_KEY_SECRET"
    if not re.fullmatch(r"rzp_(test|live)_[A-Za-z0-9]+", key_id) or any(v in (key_id + secret).lower() for v in ("placeholder", "your_key", "xxxxx")):
        return "invalid or placeholder credentials"
    mode = os.environ.get("RAZORPAY_MODE", "").strip()
    if mode and (mode not in {"test", "live"} or not key_id.startswith(f"rzp_{mode}_")):
        return "RAZORPAY_MODE does not match the key ID"
    return None


def require_payment_configuration():
    issue = payment_configuration_error()
    if issue:
        log.error("Payment configuration unavailable: %s; check backend environment and restart", issue)
        raise HTTPException(503, "Online payment is temporarily unavailable. Please contact us or try again later.")


razorpay_client = razorpay.Client(auth=(os.environ.get("RAZORPAY_KEY_ID", "").strip(), os.environ.get("RAZORPAY_KEY_SECRET", "").strip()))
razorpay_client.set_app_details({"title": "JLT Fragrances", "version": "1"})

app = FastAPI(title="JLT Fragrances API")
api = APIRouter(prefix="/api")

JWT_ALGORITHM = "HS256"


def jwt_secret() -> str:
    return os.environ["JWT_SECRET"]


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))


def create_token(email: str, role: str) -> str:
    payload = {
        "sub": email,
        "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(days=7),
        "type": "access",
    }
    return jwt.encode(payload, jwt_secret(), algorithm=JWT_ALGORITHM)


async def get_current_admin(request: Request) -> dict:
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else None
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, jwt_secret(), algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")
    if payload.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin only")
    user = await db.users.find_one({"email": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user or user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access is required")
    return user


class LoginRequest(BaseModel):
    username: Optional[str] = None
    email: Optional[str] = None
    password: str = Field(min_length=1, max_length=72)


class ContactRequest(BaseModel):
    name: str
    email: EmailStr
    phone: Optional[str] = None
    order_number: Optional[str] = None
    message: str


class QuizRequest(BaseModel):
    scent: Optional[str] = None
    occasion: Optional[str] = None
    projection: Optional[str] = None
    budget: Optional[float] = Field(default=None, gt=0, allow_inf_nan=False)
    gender: Optional[str] = None


class ReviewCreate(BaseModel):
    name: str
    rating: int = Field(ge=1, le=5)
    title: Optional[str] = None
    comment: str


class ProductUpsert(BaseModel):
    name: str
    brand_inspiration: str
    description: str = ""
    tags: List[str] = []
    category: str = "Designer"
    pricing_tier: str = "Designer/Arab"
    scent_family: List[str] = []
    moods: List[str] = []
    gender: str = "Unisex"
    occasions: List[str] = []
    seasons: List[str] = []
    longevity: str = "Long Lasting (6-8 hrs)"
    projection: str = "Moderate"
    notes: dict = Field(default_factory=lambda: {"top": [], "heart": [], "base": []})
    smells_like: str = ""
    best_for: str = ""
    who_should_buy: str = ""
    image_url: Optional[str] = None
    is_bestseller: bool = False
    is_new_arrival: bool = False
    in_stock: bool = True
    sizes: List[dict] = Field(default_factory=lambda: [
        {"size": "20ml", "price": 449},
        {"size": "50ml", "price": 1099},
        {"size": "100ml", "price": 1699},
    ])
    is_active: bool = True


class OrderItem(BaseModel):
    slug: str = Field(min_length=1, max_length=200)
    name: str = Field(default="", max_length=200)
    brand: str = Field(default="", max_length=200)
    size: str = Field(min_length=1, max_length=20)
    selections: List[str] = Field(default_factory=list, max_length=10)
    price: int = 0
    qty: int = Field(ge=1, le=20)


class CreateOrderRequest(BaseModel):
    customer_name: str = Field(min_length=2, max_length=100)
    customer_email: EmailStr
    customer_phone: str = Field(min_length=7, max_length=20)
    shipping_address: str = Field(min_length=5, max_length=400)
    shipping_city: str = Field(min_length=2, max_length=80)
    shipping_state: str = Field(min_length=2, max_length=80)
    shipping_pin: str = Field(pattern=r"^[1-9][0-9]{5}$")
    items: List[OrderItem] = Field(min_length=1, max_length=50)
    total_amount: int = 0  # ignored; the server computes the payable amount
    request_id: Optional[str] = Field(default=None, min_length=8, max_length=80)


class PaymentVerifyRequest(BaseModel):
    razorpay_order_id: str
    razorpay_payment_id: str
    razorpay_signature: str


class PricingDefaults(BaseModel):
    designer_arab: dict = Field(default_factory=lambda: {"20ml": 449, "50ml": 1099, "100ml": 1699})
    niche: dict = Field(default_factory=lambda: {"20ml": 549, "50ml": 1199, "100ml": 1799})
    ultra_niche: dict = Field(default_factory=lambda: {"20ml": 649, "50ml": 1299, "100ml": 1899})


class ShippingSettings(BaseModel):
    enabled: bool = True
    threshold: int = Field(default=2000, ge=0, le=1_000_000)
    charge_below_threshold: int = Field(default=49, ge=0, le=100_000)
    charge_at_or_above_threshold: int = Field(default=99, ge=0, le=100_000)


async def get_shipping_settings() -> dict:
    stored = await db.settings.find_one({"key": "shipping_settings"}, {"_id": 0, "value": 1})
    try:
        return ShippingSettings(**((stored or {}).get("value") or {})).model_dump()
    except Exception:
        return ShippingSettings().model_dump()


def calculate_shipping(subtotal: int, settings: dict) -> int:
    """Single source of truth for shipping, in whole rupees."""
    if not settings.get("enabled", True):
        return 0
    return settings["charge_at_or_above_threshold"] if subtotal >= settings["threshold"] else settings["charge_below_threshold"]


class ImportCommitRequest(BaseModel):
    rows: List[dict]
    mode: str = "merge"
    confirm_replace: bool = False
    legacy_mapping: bool = False


LEGACY_TIERS = {449: ("Designer", "Designer/Arab"), 549: ("Niche", "Niche"), 649: ("Ultra Niche", "Ultra Niche")}
IMPORT_HEADERS = {
    "name": "Name", "brand": "Brand Inspiration", "category": "Category", "description": "Description",
    "base_price": "Base Price", "sku": "SKU", "price_20": "20ml Price", "price_50": "50ml Price", "price_100": "100ml Price",
}


def clean_import_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


NOTE_COLUMNS = {
    "top": ("Top Notes",),
    "heart": ("Heart Notes", "Middle Notes", "Middle / Heart Notes"),
    "base": ("Base Notes",),
}


def split_note_text(value: object) -> List[str]:
    """Split one spreadsheet notes cell into its listed notes, keeping order and exact text."""
    if value is None:
        return []
    return [part.strip() for part in re.split(r"[,;]", str(value)) if part.strip()]


ATTRIBUTE_LIST_COLUMNS = {"scent_family": "Scent Family", "moods": "Moods", "occasions": "Occasions", "seasons": "Seasons", "tags": "Tags"}
ATTRIBUTE_TEXT_COLUMNS = {"gender": "Gender", "longevity": "Longevity", "projection": "Projection"}
ATTRIBUTE_FLAG_COLUMNS = {"is_bestseller": "Is Bestseller", "is_new_arrival": "Is New Arrival", "is_active": "Enabled"}
GENDER_VALUES = {"men", "women", "unisex"}


def parse_flag(value: object) -> Optional[bool]:
    text = str(value).strip().lower() if value is not None else ""
    if text in {"yes", "y", "true", "1"}: return True
    if text in {"no", "n", "false", "0"}: return False
    return None


def parse_import_attributes(get, errors: list) -> dict:
    """Read optional product attributes exactly as written in the workbook; blank cells are omitted."""
    attrs = {}
    for field, column in ATTRIBUTE_LIST_COLUMNS.items():
        items = split_note_text(get(column))
        if items: attrs[field] = items
    for field, column in ATTRIBUTE_TEXT_COLUMNS.items():
        text = str(get(column) or "").strip()
        if text: attrs[field] = text
    if "gender" in attrs and attrs["gender"].lower() not in GENDER_VALUES:
        errors.append(f"Unsupported gender: {attrs['gender']}")
    for field, column in ATTRIBUTE_FLAG_COLUMNS.items():
        raw = get(column)
        if raw is None or not str(raw).strip(): continue
        flag = parse_flag(raw)
        if flag is None: errors.append(f"Invalid {column} value: {raw}")
        else: attrs[field] = flag
    image = str(get("Image URL") or "").strip()
    if image:
        if re.match(r"^https?://", image, re.I): attrs["image_url"] = image
        else: errors.append("Image URL must start with http:// or https://")
    stock = get("Stock")
    if stock is not None and str(stock).strip():
        try:
            count = int(float(stock))
            if count < 0: raise ValueError
            attrs["stock"], attrs["in_stock"] = count, count > 0
        except (TypeError, ValueError):
            errors.append("Invalid Stock value")
    return attrs


def clean_import_attributes(value: object) -> dict:
    """Whitelist attributes returned from the preview before they reach MongoDB."""
    source = value if isinstance(value, dict) else {}
    cleaned = {}
    for field in ATTRIBUTE_LIST_COLUMNS:
        if isinstance(source.get(field), list):
            cleaned[field] = [i.strip() for i in source[field] if isinstance(i, str) and i.strip()]
    for field in ATTRIBUTE_TEXT_COLUMNS:
        if isinstance(source.get(field), str) and source[field].strip(): cleaned[field] = source[field].strip()
    for field in ATTRIBUTE_FLAG_COLUMNS:
        if isinstance(source.get(field), bool): cleaned[field] = source[field]
    if isinstance(source.get("image_url"), str) and re.match(r"^https?://", source["image_url"], re.I): cleaned["image_url"] = source["image_url"].strip()
    if isinstance(source.get("stock"), int) and not isinstance(source["stock"], bool) and source["stock"] >= 0:
        cleaned["stock"], cleaned["in_stock"] = source["stock"], source["stock"] > 0
    return cleaned


def clean_import_notes(value: object) -> dict:
    """Validate notes passed back from the import preview; never invents notes."""
    source = value if isinstance(value, dict) else {}
    cleaned = {}
    for level in NOTE_COLUMNS:
        items = source.get(level)
        cleaned[level] = [i.strip() for i in items if isinstance(i, str) and i.strip()] if isinstance(items, list) else []
    return cleaned


def parse_import_workbook(raw: bytes, legacy_mapping: bool = False) -> dict:
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(413, "Workbook exceeds the 10 MB limit")
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise HTTPException(422, f"Unable to read workbook: {exc}")
    if not values:
        raise HTTPException(422, "Workbook has no rows")
    headers = [str(v).strip() if v is not None else "" for v in values[0]]
    header_map = {name.lower(): idx for idx, name in enumerate(headers) if name}
    if "name" not in header_map or "brand inspiration" not in header_map:
        raise HTTPException(422, "Workbook requires Name and Brand Inspiration columns")
    defaults = PricingDefaults().model_dump()
    seen = {}
    rows = []
    blank_rows = []
    serial_cols = {idx for name, idx in header_map.items() if name in {"sl no", "sl. no", "s.no", "s no", "sr no", "sr. no", "#", "serial no"}}
    for source_row, raw_row in enumerate(values[1:], start=2):
        # Rows holding only a serial number are numbering placeholders, not products.
        if not any(v is not None and str(v).strip() for idx, v in enumerate(raw_row) if idx not in serial_cols):
            if any(v is not None and str(v).strip() for v in raw_row):
                blank_rows.append(source_row)
            continue
        get = lambda key: raw_row[header_map[key.lower()]] if key.lower() in header_map and header_map[key.lower()] < len(raw_row) else None
        name, brand = str(get("Name") or "").strip(), str(get("Brand Inspiration") or "").strip()
        category = str(get("Category") or "").strip()
        sku = str(get("SKU") or get("Product Code") or "").strip()
        base = get("Base Price")
        try:
            base_price = int(base) if base is not None and str(base).strip() else None
        except (TypeError, ValueError):
            base_price = None
        errors = []
        if not name: errors.append("Missing product name")
        if not brand: errors.append("Missing brand")
        if base is not None and str(base).strip() and base_price is None: errors.append("Invalid base price")
        tier_category, tier = None, None
        if category:
            normalized = clean_import_key(category)
            category_map = {"designer": ("Designer", "Designer/Arab"), "arab": ("Arab", "Designer/Arab"), "designer arab": ("Designer", "Designer/Arab"), "niche": ("Niche", "Niche"), "ultra niche": ("Ultra Niche", "Ultra Niche")}
            if normalized not in category_map: errors.append(f"Unsupported category: {category}")
            else: tier_category, tier = category_map[normalized]
        elif legacy_mapping and base_price in LEGACY_TIERS:
            tier_category, tier = LEGACY_TIERS[base_price]
        else:
            errors.append("Category missing and base price cannot resolve a tier")
        tier_key = {"Designer/Arab": "designer_arab", "Niche": "niche", "Ultra Niche": "ultra_niche"}.get(tier)
        prices = {}
        for size, column in (("20ml", "20ml Price"), ("50ml", "50ml Price"), ("100ml", "100ml Price")):
            explicit = get(column)
            if explicit is not None and str(explicit).strip():
                try: prices[size] = int(explicit)
                except (TypeError, ValueError): errors.append(f"Invalid {size} price")
                if size in prices and prices[size] <= 0:
                    errors.append(f"Invalid {size} price: must be positive")
        if tier_key:
            prices = {size: prices.get(size, defaults[tier_key][size]) for size in ("20ml", "50ml", "100ml")}
        notes = {}
        for level, columns in NOTE_COLUMNS.items():
            column = next((c for c in columns if c.lower() in header_map), None)
            notes[level] = split_note_text(get(column)) if column else []
        attributes = parse_import_attributes(get, errors)
        key = sku.lower() if sku else f"{clean_import_key(brand)}::{clean_import_key(name)}"
        fingerprint = (str(get("Description") or "").strip(), json.dumps(notes, sort_keys=True), json.dumps(attributes, sort_keys=True, default=str))
        first = seen.get(key)
        if first:
            errors.append(f"Duplicate spreadsheet identifier: same brand + name as row {first['row']}" + ("; identical content" if first["fp"] == fingerprint else "; content differs - add a SKU column to keep both"))
        else:
            seen[key] = {"row": source_row, "fp": fingerprint}
        rows.append({"source_row": source_row, "name": name, "brand": brand, "category": tier_category or category, "pricing_tier": tier, "sku": sku, "matching_identifier": key, "description": str(get("Description") or "").strip(), "notes": notes, "attributes": attributes, "sizes": [{"size": s, "price": prices[s]} for s in ("20ml", "50ml", "100ml")] if len(prices) == 3 else [], "action": "REJECT" if errors else "ADD", "validation_message": "; ".join(errors), "legacy_mapping": not bool(category) and bool(tier)} )
    return {"headers": headers, "rows": rows, "total_rows": len(rows), "valid_rows": sum(r["action"] != "REJECT" for r in rows), "rejected_rows": sum(r["action"] == "REJECT" for r in rows), "duplicate_rows": sum("Duplicate spreadsheet identifier" in r["validation_message"] for r in rows), "blank_rows": blank_rows}


def validate_pricing_defaults(value: dict) -> dict:
    required = {"20ml", "50ml", "100ml"}
    for tier in ("designer_arab", "niche", "ultra_niche"):
        prices = value.get(tier, {})
        if set(prices) != required or any(not isinstance(prices[size], int) or prices[size] <= 0 for size in required):
            raise HTTPException(422, f"Invalid pricing defaults for {tier}")
    return value


# Razorpay helper functions
def verify_razorpay_signature(order_id: str, payment_id: str, signature: str) -> bool:
    """Verify Razorpay payment signature."""
    try:
        message = f"{order_id}|{payment_id}"
        secret = os.environ.get("RAZORPAY_KEY_SECRET", "").strip().encode()
        computed_sig = hmac.new(secret, message.encode(), hashlib.sha256).hexdigest()
        return bool(secret) and hmac.compare_digest(computed_sig, signature)
    except Exception as e:
        log.error(f"Signature verification error: {e}")
        return False


async def send_telegram_notification(order_data: dict) -> bool:
    """Send Telegram notification. Return True if successful, False otherwise."""
    try:
        token = os.environ.get("TELEGRAM_BOT_TOKEN", "")
        chat_id = os.environ.get("TELEGRAM_CHAT_ID", "")
        
        if not token or not chat_id:
            log.warning("Telegram credentials not configured, skipping notification")
            return False
        
        import aiohttp
        
        # Format message
        msg = f"""
✅ New Order Confirmed!

📋 Order ID: {order_data.get('order_id')}
💳 Razorpay Payment ID: {order_data.get('razorpay_payment_id')}

👤 Customer: {order_data.get('customer_name')}
📧 Email: {order_data.get('customer_email')}
📞 Phone: {order_data.get('customer_phone')}

📦 Shipping:
{order_data.get('shipping_address')}
{order_data.get('shipping_city')}, {order_data.get('shipping_state')} {order_data.get('shipping_pin')}

🛍️ Items:
"""
        for item in order_data.get('items', []):
            msg += f"\n  • {item['name']} ({item['brand']}) {item['size']} x{item['qty']} = ₹{item['price'] * item['qty']}"
        
        msg += f"\n\n💰 Total: ₹{order_data.get('total_amount') / 100:.2f}"
        msg += f"\n✔️ Payment Status: PAID"
        msg += f"\n⏰ Timestamp: {order_data.get('created_at')}"
        
        # Send via Telegram
        url = f"https://api.telegram.org/bot{token}/sendMessage"
        payload = {"chat_id": chat_id, "text": msg}
        
        async with aiohttp.ClientSession() as session:
            async with session.post(url, json=payload, timeout=aiohttp.ClientTimeout(total=5)) as resp:
                if resp.status == 200:
                    log.info(f"Telegram notification sent for order {order_data.get('order_id')}")
                    return True
                else:
                    log.error(f"Telegram API error: {resp.status}")
                    return False
    except Exception as e:
        log.error(f"Telegram notification error (non-blocking): {e}")
        return False  # Don't break checkout if Telegram fails


def format_inr(amount_paise: int) -> str:
    return f"₹{amount_paise / 100:,.2f}"


def send_admin_order_email(order_data: dict) -> bool:
    """Send admin order notification email for confirmed orders."""
    smtp_host = os.environ.get("SMTP_HOST", "").strip()
    smtp_port = os.environ.get("SMTP_PORT", "").strip()
    smtp_user = os.environ.get("SMTP_USER", "").strip()
    smtp_password = os.environ.get("SMTP_PASSWORD", "").strip()
    email_from = os.environ.get("EMAIL_FROM", "").strip()
    email_to = os.environ.get("EMAIL_TO", "").strip()

    if not (smtp_host and smtp_port and smtp_user and smtp_password and email_from and email_to):
        log.error("SMTP configuration is incomplete, admin order email skipped")
        return False

    try:
        port = int(smtp_port)
    except ValueError:
        log.error("SMTP_PORT must be an integer, admin order email skipped")
        return False

    subject = f"New Order Received - {order_data.get('order_id')}"
    body_lines = [
        f"Order ID: {order_data.get('order_id')}",
        f"Date & Time: {order_data.get('created_at')}",
        f"Customer Name: {order_data.get('customer_name')}",
        f"Customer Email: {order_data.get('customer_email')}",
        f"Customer Phone: {order_data.get('customer_phone')}",
        "",
        "Shipping Address:",
        f"{order_data.get('shipping_address')}",
        f"{order_data.get('shipping_city')}, {order_data.get('shipping_state')} {order_data.get('shipping_pin')}",
        "",
        f"Payment Method: {order_data.get('payment_method', 'Razorpay')}",
        f"Payment Status: {order_data.get('payment_status', 'paid').upper()}",
        f"Razorpay Payment ID: {order_data.get('razorpay_payment_id') or 'N/A'}",
        f"Razorpay Order ID: {order_data.get('razorpay_order_id') or 'N/A'}",
        "",
        "Items:",
    ]

    for item in order_data.get("items", []):
        unit_price = item.get("price", 0)
        qty = item.get("qty", 0)
        line_total = unit_price * qty
        body_lines.append(
            f"- {item.get('name')} | Brand: {item.get('brand')} | Size: {item.get('size')} | Quantity: {qty} | Unit Price: {format_inr(unit_price)} | Line Total: {format_inr(line_total)}"
        )

    body_lines.extend([
        "",
        f"Grand Total: {format_inr(order_data.get('total_amount', 0))}",
    ])

    message = EmailMessage()
    message["From"] = email_from
    message["To"] = email_to
    message["Subject"] = subject
    message.set_content("\n".join(body_lines))

    try:
        if port == 465:
            server = smtplib.SMTP_SSL(smtp_host, port, timeout=10)
        else:
            server = smtplib.SMTP(smtp_host, port, timeout=10)
            server.starttls()

        with server:
            server.login(smtp_user, smtp_password)
            server.send_message(message)

        log.info(
            "Admin order email sent",
            extra={
                "order_id": order_data.get("order_id"),
                "recipient": email_to,
            },
        )
        return True
    except Exception as e:
        log.error(
            "Admin order email failed",
            extra={
                "order_id": order_data.get("order_id"),
                "error": str(e),
            },
        )
        return False


async def send_admin_order_email_async(order_data: dict) -> bool:
    import asyncio
    return await asyncio.to_thread(send_admin_order_email, order_data)


@api.get("/")
async def root():
    return {"message": "JLT Fragrances API", "products": await db.products.count_documents({"is_active": {"$ne": False}})}


@api.get("/products")
async def list_products(
    q: Optional[str] = None,
    brand: Optional[List[str]] = Query(None), category: Optional[List[str]] = Query(None),
    tier: Optional[List[str]] = Query(None), scent: Optional[List[str]] = Query(None),
    note: Optional[List[str]] = Query(None), occasion: Optional[List[str]] = Query(None),
    season: Optional[List[str]] = Query(None), gender: Optional[List[str]] = Query(None),
    mood: Optional[List[str]] = Query(None), size: Optional[List[str]] = Query(None),
    longevity: Optional[List[str]] = Query(None), projection: Optional[List[str]] = Query(None),
    bestseller: Optional[bool] = None, new_arrival: Optional[bool] = None, niche: Optional[bool] = None,
    min_price: Optional[int] = Query(None, ge=0), max_price: Optional[int] = Query(None, ge=0),
    sort: str = "popular", page: int = Query(1, ge=1), limit: int = Query(24, ge=1, le=120),
):
    if min_price is not None and max_price is not None and min_price > max_price:
        raise HTTPException(422, "Minimum price must not exceed maximum price")
    selected = {"brand":brand or [], "category":category or [], "tier":tier or [], "scent":scent or [],
        "note":note or [], "occasion":occasion or [], "season":season or [], "gender":gender or [],
        "mood":mood or [], "size":size or [], "longevity":longevity or [], "projection":projection or []}
    products = [normalized(p) for p in await db.products.find({"is_active": {"$ne": False}}, {"_id":0}).to_list(None)]
    items = [p for p in products if matches(p,selected,q,min_price,max_price,bestseller,new_arrival,niche)]
    def available_price(p):
        variants = [s["price"] for s in p["sizes"] if not size or key(s["size"]) in {key(v) for v in size}]
        return min(variants, default=0)
    if sort in ("price_asc","price_desc"):
        items.sort(key=lambda p: (available_price(p), key(p["name"])), reverse=sort=="price_desc")
    elif sort=="rating":
        items.sort(key=lambda p: (-(p.get("rating") or 0), key(p["name"])))
    elif sort=="new":
        items.sort(key=lambda p: (not p.get("is_new_arrival"), key(p["name"])))
    else:
        items.sort(key=lambda p: (not p.get("is_bestseller"), -(p.get("rating") or 0), key(p["name"]),p["slug"]))
    total=len(items)
    return {"items":items[(page-1)*limit:page*limit], "total":total,"page":page,"limit":limit}


@api.get("/products/filters")
async def get_filters():
    products = [normalized(p) for p in await db.products.find({"is_active": {"$ne": False}}, {"_id":0}).to_list(None)]
    return filter_options(products)


@api.get("/brands")
async def list_brands():
    """Alphabetical brand directory with product counts."""
    pipeline = [
        {"$match": {"is_active": {"$ne": False}}},
        {"$group": {
            "_id": "$brand_inspiration",
            "count": {"$sum": 1},
            "is_niche": {"$first": "$is_niche"},
        }},
        {"$sort": {"_id": 1}},
    ]
    cursor = db.products.aggregate(pipeline)
    items = [{"brand": d["_id"], "count": d["count"], "is_niche": bool(d.get("is_niche"))} async for d in cursor]
    grouped: dict = {}
    for it in items:
        letter = (it["brand"][:1] or "#").upper()
        grouped.setdefault(letter, []).append(it)
    return {"groups": [{"letter": k, "brands": v} for k, v in sorted(grouped.items())], "total": len(items)}


async def offer_catalogue():
    combos, sets = copy.deepcopy(COMBOS), copy.deepcopy(DISCOVERY_SETS)
    saved = await db.settings.find_one({"key": "offer_prices"}, {"_id": 0}) or {}
    for offer in [*combos["tester_combos"], *combos["special_offers"], *sets]:
        price = saved.get("prices", {}).get(offer["slug"])
        if type(price) is int and price > 0:
            offer["price"] = price
    return combos, sets


@api.get("/combos")
async def list_combos():
    return (await offer_catalogue())[0]


@api.get("/admin/offer-prices")
async def get_offer_prices(admin: dict = Depends(get_current_admin)):
    combos, sets = await offer_catalogue()
    return {"items": [*combos["tester_combos"], *combos["special_offers"], *sets]}


@api.put("/admin/offer-prices")
async def save_offer_prices(body: dict, admin: dict = Depends(get_current_admin)):
    allowed = {o["slug"] for o in [*COMBOS["tester_combos"], *COMBOS["special_offers"], *DISCOVERY_SETS]}
    prices = body.get("prices")
    if not isinstance(prices, dict) or set(prices) != allowed or any(type(v) is not int or not 0 < v <= 1_000_000 for v in prices.values()):
        raise HTTPException(422, "Provide a positive whole-rupee price for every current offer (maximum ₹1,000,000).")
    await db.settings.update_one({"key": "offer_prices"}, {"$set": {"prices": prices}}, upsert=True)
    return await get_offer_prices(admin)


@api.get("/products/bestsellers")
async def bestsellers(limit: int = 12):
    items = await db.products.find({"is_bestseller": True, "is_active": {"$ne": False}}, {"_id": 0}).limit(limit).to_list(limit)
    if not items:
        items = await db.products.find({"is_active": {"$ne": False}}, {"_id": 0}).sort([("rating", -1), ("name", 1)]).limit(limit).to_list(limit)
    return {"items": items}


@api.get("/products/{slug}")
async def get_product(slug: str):
    p = await db.products.find_one({"slug": slug, "is_active": {"$ne": False}}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404, detail="Product not found")
    p = normalized(p)
    related_q = {
        "is_active": {"$ne": False},
        "slug": {"$ne": slug},
        "$or": [
            {"brand_inspiration": p["brand_inspiration"]},
            {"scent_family": {"$in": p.get("scent_family", [])}},
        ],
    }
    related = await db.products.find(related_q, {"_id": 0}).limit(6).to_list(6)
    return {"product": p, "related": related}


async def quiz_inventory():
    return inventory(await db.products.find({"is_active": {"$ne": False}, "in_stock": {"$ne": False}}, {"_id": 0}).to_list(None))


@api.post("/quiz/options")
async def quiz_options(body: QuizRequest):
    answers = body.model_dump(exclude_none=True)
    prefix = {}
    for step in STEPS:
        if step not in answers:
            remaining = matching(await quiz_inventory(), prefix)
            return {"step": step, "options": choices(remaining, step), "total": len(remaining)}
        prefix[step] = answers[step]
    raise HTTPException(422, "All questions are answered; request recommendations.")


@api.post("/quiz/recommend")
async def quiz_recommend(body: QuizRequest):
    answers = body.model_dump(exclude_none=True)
    if set(answers) != set(STEPS):
        raise HTTPException(422, "Complete all five quiz questions.")
    items = matching(await quiz_inventory(), answers)
    items.sort(key=lambda p: (-(p.get("rating") or 0), not p.get("is_bestseller")))
    return {"items": items[:12], "total": len(items)}


@api.get("/discovery-sets")
async def discovery_sets():
    return {"items": (await offer_catalogue())[1]}


@api.post("/contact")
async def contact(body: ContactRequest):
    doc = body.model_dump()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    doc["id"] = str(uuid.uuid4())
    await db.contact_messages.insert_one(doc)
    return {"ok": True, "message": "Thanks! We'll get back to you within 24 hours."}


@api.get("/reviews/{slug}")
async def get_reviews(slug: str):
    items = await db.reviews.find({"product_slug": slug}, {"_id": 0}).sort([("created_at", -1)]).to_list(50)
    return {"items": items}


@api.post("/reviews/{slug}")
async def add_review(slug: str, body: ReviewCreate):
    doc = body.model_dump()
    doc.update({"id": str(uuid.uuid4()), "product_slug": slug, "created_at": datetime.now(timezone.utc).isoformat()})
    await db.reviews.insert_one(doc)
    doc.pop("_id", None)
    return {"ok": True, "review": doc}


@api.post("/auth/login")
async def login(body: LoginRequest):
    identifier = (body.username or body.email or "").strip().lower()
    user = await db.users.find_one({"$or": [{"email": identifier}, {"username": identifier}]})
    if not user or user.get("role") != "admin" or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid username or password")
    token = create_token(user["email"], user.get("role", "admin"))
    return {"token": token, "user": {"email": user["email"], "name": user.get("name"), "role": user.get("role")}}


@api.get("/auth/me")
async def me(admin: dict = Depends(get_current_admin)):
    return admin


@api.post("/admin/products")
async def create_product(body: ProductUpsert, admin: dict = Depends(get_current_admin)):
    cnt = await db.products.count_documents({})
    p = body.model_dump()
    p["slug"] = slugify(f"{p['name']}-inspired-by-{p['brand_inspiration']}-{cnt + 1}")
    size_values = {item["size"]: item["price"] for item in p.get("sizes", []) if item.get("size") in {"20ml", "50ml", "100ml"}}
    if len(size_values) != len(p.get("sizes", [])) or len(size_values) != 3 or any(not isinstance(v, int) or v <= 0 for v in size_values.values()):
        raise HTTPException(422, "Exactly one valid price is required for each 20ml, 50ml, and 100ml variant")
    p["sizes"] = [{"size": size, "price": size_values[size]} for size in ("20ml", "50ml", "100ml")]
    p["base_price"] = p["sizes"][0]["price"]
    p["rating"] = 0
    p["review_count"] = 0
    await db.products.insert_one(p)
    p.pop("_id", None)
    return {"ok": True, "product": p}


@api.put("/admin/products/{slug}")
async def update_product(slug: str, body: ProductUpsert, admin: dict = Depends(get_current_admin)):
    existing = await db.products.find_one({"slug": slug})
    if not existing:
        raise HTTPException(404, "Product not found")
    values = body.model_dump()
    size_values = {item["size"]: item["price"] for item in values.get("sizes", []) if item.get("size") in {"20ml", "50ml", "100ml"}}
    if len(size_values) != 3 or len(size_values) != len(values.get("sizes", [])) or any(not isinstance(v, int) or v <= 0 for v in size_values.values()):
        raise HTTPException(422, "Exactly one valid price is required for each 20ml, 50ml, and 100ml variant")
    values["sizes"] = [{"size": size, "price": size_values[size]} for size in ("20ml", "50ml", "100ml")]
    values["base_price"] = values["sizes"][0]["price"]
    await db.products.update_one({"slug": slug}, {"$set": values})
    existing.pop("_id", None)
    return {"ok": True, "product": {**existing, **values, "slug": slug}}


@api.delete("/admin/products/{slug}")
async def delete_product(slug: str, admin: dict = Depends(get_current_admin)):
    res = await db.products.update_one({"slug": slug}, {"$set": {"is_active": False, "archived_at": datetime.now(timezone.utc).isoformat()}})
    if res.matched_count == 0:
        raise HTTPException(404, "Not found")
    return {"ok": True, "archived": True}


@api.post("/admin/products/import-csv")
async def import_csv(file: UploadFile = File(...), admin: dict = Depends(get_current_admin)):
    content = (await file.read()).decode("utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(content))
    inserted = 0
    seen = set()
    for row in reader:
        brand = (row.get("brand_inspiration") or row.get("brand") or "").strip()
        name = (row.get("name") or row.get("product_name") or "").strip()
        if not brand or not name:
            continue
        stable_slug = slugify(f"{name}-inspired-by-{brand}")
        if stable_slug in seen:
            continue
        seen.add(stable_slug)
        p = build_product(1, brand, name)
        p["slug"] = stable_slug
        p["is_active"] = True
        await db.products.update_one({"slug": p["slug"]}, {"$set": p}, upsert=True)
        inserted += 1
    return {"ok": True, "inserted": inserted}


@api.get("/admin/pricing-defaults")
async def get_pricing_defaults(admin: dict = Depends(get_current_admin)):
    defaults = await db.settings.find_one({"key": "pricing_defaults"}, {"_id": 0, "value": 1})
    return defaults.get("value", PricingDefaults().model_dump()) if defaults else PricingDefaults().model_dump()


@api.put("/admin/pricing-defaults")
async def update_pricing_defaults(body: PricingDefaults, apply_to_products: bool = False, admin: dict = Depends(get_current_admin)):
    value = validate_pricing_defaults(body.model_dump())
    await db.settings.update_one({"key": "pricing_defaults"}, {"$set": {"key": "pricing_defaults", "value": value, "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    if apply_to_products:
        tier_map = {"Designer/Arab": "designer_arab", "Niche": "niche", "Ultra Niche": "ultra_niche"}
        for tier, key in tier_map.items():
            prices = value[key]
            await db.products.update_many({"pricing_tier": tier, "is_active": {"$ne": False}}, {"$set": {"sizes": [{"size": size, "price": prices[size]} for size in ("20ml", "50ml", "100ml")], "base_price": prices["20ml"]}})
    return {"ok": True, "defaults": value, "applied": apply_to_products}


@api.get("/shipping-settings")
async def public_shipping_settings():
    return await get_shipping_settings()


@api.get("/admin/shipping-settings")
async def admin_get_shipping_settings(admin: dict = Depends(get_current_admin)):
    return await get_shipping_settings()


@api.put("/admin/shipping-settings")
async def admin_update_shipping_settings(body: ShippingSettings, admin: dict = Depends(get_current_admin)):
    value = body.model_dump()
    await db.settings.update_one({"key": "shipping_settings"}, {"$set": {"key": "shipping_settings", "value": value, "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    return {"ok": True, "settings": value}


class SeoEntry(BaseModel):
    path: str = Field(min_length=1, max_length=300, pattern=r"^/[A-Za-z0-9\-_/.]*$")
    title: str = Field(default="", max_length=120)
    description: str = Field(default="", max_length=320)
    canonical_url: str = Field(default="", max_length=500, pattern=r"^(https?://\S+)?$")
    og_title: str = Field(default="", max_length=120)
    og_description: str = Field(default="", max_length=320)
    og_image: str = Field(default="", max_length=500, pattern=r"^(https?://\S+|/\S*)?$")
    noindex: bool = False


SITE_NAME = "JLT Fragrances"
STATIC_SEO = {
    "/": ("JLT Fragrances – Premium Inspired Perfumes Online India", "Shop 800+ Luxury Inspired Scents in 20ml, 50ml and 100ml. Designer, Arab and niche inspired scents with pan-India delivery."),
    "/shop": ("Shop Perfumes – Designer, Arab & Niche Inspired", "Browse and filter the full JLT Fragrances catalogue by brand, notes, occasion and size."),
    "/brands": ("Fragrances by Brand Inspiration", "Explore perfumes inspired by your favourite fragrance houses."),
    "/bestsellers": ("Bestselling Perfumes", "Our most loved premium-inspired fragrances."),
    "/discovery-sets": ("Perfume Discovery Sets", "Try more scents with JLT Fragrances discovery sets and bundles."),
    "/find-your-scent": ("Find Your Scent – Fragrance Quiz", "Answer a few questions and get fragrance recommendations matched to your taste."),
    "/about": ("About JLT Fragrances", "Learn about JLT Fragrances and our approach to premium-inspired perfumes."),
    "/contact": ("Contact JLT Fragrances", "Get in touch with JLT Fragrances for order support and recommendations."),
    "/shipping-policy": ("Shipping Policy", "Shipping charges and delivery information for JLT Fragrances orders."),
    "/refund-policy": ("Refund Policy", "Returns and refund information for JLT Fragrances."),
    "/privacy-policy": ("Privacy Policy", "How JLT Fragrances collects and uses your information."),
    "/terms": ("Terms & Conditions", "Terms of use for the JLT Fragrances website."),
}
NOINDEX_PATHS = ("/admin", "/cart", "/checkout", "/wishlist")


def site_url() -> str:
    return os.environ.get("SITE_URL", "").rstrip("/")


def clip(text: str, limit: int) -> str:
    text = " ".join((text or "").split())
    return text if len(text) <= limit else text[: limit - 1].rsplit(" ", 1)[0] + "…"


async def build_seo(path: str) -> dict:
    """Admin overrides win; otherwise defaults come from real page/product data."""
    path = "/" + path.strip("/") if path.strip("/") else "/"
    seo = {"path": path, "title": "", "description": "", "canonical_url": "", "og_title": "", "og_description": "", "og_image": "", "noindex": path.startswith(NOINDEX_PATHS)}
    if path in STATIC_SEO:
        seo["title"], seo["description"] = STATIC_SEO[path]
    elif path.startswith("/product/"):
        product = await db.products.find_one({"slug": path[len("/product/"):], "is_active": {"$ne": False}}, {"_id": 0})
        if product:
            p = normalized(product)
            seo["title"] = f"{p['name']} – Inspired by {p['brand_inspiration']}"
            notes = ", ".join(dict.fromkeys(n for part in p["notes"].values() for n in part))
            seo["description"] = clip(p.get("description") or f"{p['name']} inspired by {p['brand_inspiration']}. Notes: {notes}.", 160)
            seo["og_image"] = p.get("image_url") or ""
        else:
            seo["noindex"] = True
    elif path.startswith("/brand/"):
        seo["title"] = f"Perfumes inspired by {path[len('/brand/'):].replace('-', ' ').title()}"
    override = await db.seo.find_one({"path": path}, {"_id": 0})
    if override:
        seo.update({k: v for k, v in override.items() if k in seo and v not in ("", None)})
        seo["noindex"] = bool(override.get("noindex", seo["noindex"]))
    seo["site_name"] = SITE_NAME
    for field in ("title", "description", "og_title", "og_description"):
        seo[field] = re.sub(r"750\s*\+\s*premium[\s-]*inspired\s*(?:scents|perfumes)", "800+ Luxury Inspired Scents", seo[field], flags=re.I)
    seo["title"] = seo["title"] if SITE_NAME in seo["title"] else (f"{seo['title']} | {SITE_NAME}" if seo["title"] else SITE_NAME)
    return seo


@api.get("/seo")
async def public_seo(path: str = "/"):
    return await build_seo(path)


@api.get("/admin/seo")
async def admin_list_seo(admin: dict = Depends(get_current_admin)):
    items = await db.seo.find({}, {"_id": 0}).sort("path", 1).to_list(1000)
    return {"items": items, "static_paths": list(STATIC_SEO)}


@api.put("/admin/seo")
async def admin_upsert_seo(body: SeoEntry, admin: dict = Depends(get_current_admin)):
    value = body.model_dump()
    await db.seo.update_one({"path": body.path}, {"$set": {**value, "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    return {"ok": True, "seo": value}


@api.delete("/admin/seo")
async def admin_delete_seo(path: str, admin: dict = Depends(get_current_admin)):
    result = await db.seo.delete_one({"path": path})
    if result.deleted_count == 0:
        raise HTTPException(404, "No custom SEO for this path")
    return {"ok": True}


@app.get("/robots.txt", include_in_schema=False)
async def robots_txt():
    lines = ["User-agent: *", "Allow: /", "Disallow: /admin", "Disallow: /cart", "Disallow: /checkout", "Disallow: /wishlist", "Disallow: /api/"]
    if site_url():
        lines.append(f"Sitemap: {site_url()}/sitemap.xml")
    return PlainTextResponse("\n".join(lines) + "\n")


@app.get("/sitemap.xml", include_in_schema=False)
async def sitemap_xml():
    base = site_url()
    noindexed = {e["path"] async for e in db.seo.find({"noindex": True}, {"_id": 0, "path": 1})}
    paths = list(STATIC_SEO)
    brands = sorted({p["brand_inspiration"] async for p in db.products.find({"is_active": {"$ne": False}}, {"_id": 0, "brand_inspiration": 1})})
    paths += [f"/brand/{slugify(b)}" for b in brands]
    async for p in db.products.find({"is_active": {"$ne": False}}, {"_id": 0, "slug": 1, "updated_at": 1}):
        paths.append(f"/product/{p['slug']}")
    body = "".join(f"<url><loc>{xml_escape(base + path)}</loc></url>" for path in paths if path not in noindexed)
    return Response(f'<?xml version="1.0" encoding="UTF-8"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">{body}</urlset>', media_type="application/xml")


@api.post("/admin/import/preview")
async def preview_excel_import(file: UploadFile = File(...), legacy_mapping: bool = False, admin: dict = Depends(get_current_admin)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(415, "Only .xlsx files are supported")
    result = parse_import_workbook(await file.read(), legacy_mapping=legacy_mapping)
    for row in result["rows"]:
        existing = await db.products.find_one({"import_identifier": row["matching_identifier"]}, {"_id": 0, "slug": 1})
        if row["action"] != "REJECT": row["action"] = "UPDATE" if existing else "ADD"
        row["existing_slug"] = existing.get("slug") if existing else None
    return result


@api.post("/admin/import/commit")
async def commit_excel_import(body: ImportCommitRequest, admin: dict = Depends(get_current_admin)):
    if body.mode not in {"merge", "replace"}:
        raise HTTPException(422, "Mode must be merge or replace")
    if body.mode == "replace" and not body.confirm_replace:
        raise HTTPException(400, "Replace requires explicit confirmation")
    valid = [row for row in body.rows if row.get("action") != "REJECT"]
    identifiers = [row.get("matching_identifier") for row in valid]
    if not identifiers or len(identifiers) != len(set(identifiers)):
        raise HTTPException(422, "Duplicate matching identifiers")
    now = datetime.now(timezone.utc).isoformat()
    docs = []
    for row in valid:
        if not row.get("name") or not row.get("brand") or len(row.get("sizes", [])) != 3:
            raise HTTPException(422, f"Invalid row {row.get('source_row')}")
        slug = slugify(f"{row['name']}-inspired-by-{row['brand']}")
        doc = build_product(1, row["brand"], row["name"])
        managed = {"name": row["name"], "brand_inspiration": row["brand"], "category": row.get("category"), "pricing_tier": row.get("pricing_tier"), "description": row.get("description", ""), "notes": clean_import_notes(row.get("notes")), "sizes": row["sizes"], "base_price": row["sizes"][0]["price"], "import_identifier": row["matching_identifier"], "is_active": True, **clean_import_attributes(row.get("attributes"))}
        doc.update({"slug": slug, **managed})
        docs.append({"doc": doc, "managed": managed})
    result = {"added": 0, "updated": 0, "unchanged": 0, "archived": 0, "rejected": len(body.rows) - len(valid), "duplicates": sum("Duplicate spreadsheet identifier" in (r.get("validation_message") or "") for r in body.rows)}
    try:
        for entry in docs:
            doc, managed = entry["doc"], entry["managed"]
            old = await db.products.find_one({"import_identifier": managed["import_identifier"]})
            if old:
                if all(old.get(k) == v for k, v in managed.items()): result["unchanged"] += 1
                else: result["updated"] += 1
            else: result["added"] += 1
            # Existing products keep their slug, images, ratings and reviews; only Excel-supplied fields are written.
            defaults = {k: v for k, v in doc.items() if k not in managed}
            await db.products.update_one({"import_identifier": managed["import_identifier"]}, {"$set": {**managed, "updated_at": now}, "$setOnInsert": defaults}, upsert=True)
        if body.mode == "replace":
            keep = [entry["managed"]["import_identifier"] for entry in docs]
            archived = await db.products.update_many({"is_active": {"$ne": False}, "import_identifier": {"$nin": keep}}, {"$set": {"is_active": False, "archived_at": now}})
            result["archived"] = archived.modified_count
    except Exception as exc:
        raise HTTPException(500, f"Import failed before completion: {exc}")
    return {"ok": True, **result, "total_source_rows": len(body.rows)}


@api.get("/admin/contacts")
async def list_contacts(admin: dict = Depends(get_current_admin)):
    items = await db.contact_messages.find({}, {"_id": 0}).sort([("created_at", -1)]).limit(200).to_list(200)
    return {"items": items}


# ============ RAZORPAY & ORDER ENDPOINTS ============

async def price_cart(items):
    combos, sets = await offer_catalogue()
    offers = {o["slug"]: o for o in [*combos["tester_combos"], *combos["special_offers"], *sets]}
    authoritative_items = []
    subtotal_amount = 0
    for item in items:
        if item.slug == "special-4-30ml":
            raise HTTPException(400, "The 30ml bundle is no longer available. Remove it from your cart.")
        offer = offers.get(item.slug)
        if offer:
            if item.size != offer["size"]:
                raise HTTPException(400, "Invalid offer size")
            selected = []
            if item.selections or not offer.get("curated"):
                if len(item.selections) != offer["items"]:
                    raise HTTPException(400, f"Choose {offer['items']} fragrances for {offer['title']}")
                for slug in item.selections:
                    product = await db.products.find_one({"slug": slug, "is_active": {"$ne": False}, "in_stock": {"$ne": False}})
                    if not product or (item.size != "8ml" and not any(v.get("size") == item.size and v.get("in_stock") is not False for v in product.get("sizes", []))):
                        raise HTTPException(400, "A selected fragrance is unavailable. Please choose again.")
                    selected.append({"slug": slug, "name": product["name"], "brand": product["brand_inspiration"]})
            line = {"slug": item.slug, "name": offer["title"], "brand": "JLT Fragrances", "size": offer["size"], "price": offer["price"], "qty": item.qty,
                    "selections": item.selections, "selected_products": selected}
        else:
            product = await db.products.find_one({"slug": item.slug, "is_active": {"$ne": False}})
            if not product:
                raise HTTPException(400, f"Product {item.slug} not found")
            variant = next((v for v in product.get("sizes", []) if v.get("size") == item.size and v.get("in_stock") is not False), None)
            if not variant or not isinstance(variant.get("price"), (int, float)) or variant["price"] <= 0:
                raise HTTPException(400, f"Size {item.size} is not available for {item.slug}")
            if product.get("in_stock") is False:
                raise HTTPException(400, f"{product['name']} is currently out of stock")
            line = {"slug": item.slug, "name": product["name"], "brand": product["brand_inspiration"], "size": item.size, "price": int(variant["price"]), "qty": item.qty}
        subtotal_amount += line["price"] * line["qty"] * 100
        authoritative_items.append(line)
    shipping_amount = calculate_shipping(subtotal_amount // 100, await get_shipping_settings()) * 100
    return {"items": authoritative_items, "subtotal_amount": subtotal_amount, "shipping_amount": shipping_amount, "total_amount": subtotal_amount + shipping_amount}


class CartQuoteRequest(BaseModel):
    items: List[OrderItem] = Field(min_length=1, max_length=50)


@api.post("/cart/quote")
async def quote_cart(body: CartQuoteRequest):
    return await price_cart(body.items)


@api.post("/orders")
async def create_order(body: CreateOrderRequest):
    """Create order and Razorpay payment."""
    try:
        require_payment_configuration()
        quote = await price_cart(body.items)
        authoritative_items = quote["items"]
        subtotal_amount, shipping_amount, total_amount = (quote[k] for k in ("subtotal_amount", "shipping_amount", "total_amount"))

        # A retry with the same request_id and an identical cart reuses the existing order instead of creating another.
        if body.request_id:
            existing = await db.orders.find_one({"request_id": body.request_id, "customer_email": body.customer_email}, {"_id": 0})
            if existing and existing.get("items") == authoritative_items and existing.get("total_amount") == total_amount:
                return {
                    "ok": True, "reused": True, "payment_status": existing.get("payment_status"),
                    "order_id": existing["order_id"], "razorpay_order_id": existing["razorpay_order_id"],
                    "razorpay_key_id": os.environ.get("RAZORPAY_KEY_ID", ""),
                    "customer_name": existing["customer_name"], "customer_email": existing["customer_email"],
                    "subtotal_amount": existing.get("subtotal_amount", subtotal_amount), "shipping_amount": existing.get("shipping_amount", shipping_amount),
                    "total_amount": existing["total_amount"],
                }
        
        # Create order ID
        order_id = f"ORD-{datetime.now(timezone.utc).strftime('%Y%m%d')}-{str(uuid.uuid4())[:8].upper()}"
        
        # Create Razorpay order (amount in paise)
        razorpay_order = await run_in_threadpool(razorpay_client.order.create,
            data={"amount": total_amount, "currency": "INR", "receipt": order_id}, timeout=20
        )
        
        # Store order in MongoDB (status: pending until payment verified)
        order_doc = {
            "order_id": order_id,
            "razorpay_order_id": razorpay_order["id"],
            "customer_name": body.customer_name,
            "customer_email": body.customer_email,
            "customer_phone": body.customer_phone,
            "shipping_address": body.shipping_address,
            "shipping_city": body.shipping_city,
            "shipping_state": body.shipping_state,
            "shipping_pin": body.shipping_pin,
            "items": authoritative_items,
            "subtotal_amount": subtotal_amount,
            "shipping_amount": shipping_amount,
            "total_amount": total_amount,
            "request_id": body.request_id,
            "payment_status": "pending",
            "razorpay_payment_id": None,
            "razorpay_signature": None,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.orders.insert_one(order_doc)
        log.info(
            "Order created",
            extra={
                "order_id": order_id,
                "razorpay_order_id": razorpay_order["id"],
                "customer_email": body.customer_email,
                "total_amount": total_amount,
            },
        )
        
        return {
            "ok": True,
            "order_id": order_id,
            "razorpay_order_id": razorpay_order["id"],
            "razorpay_key_id": os.environ.get("RAZORPAY_KEY_ID", ""),
            "customer_name": body.customer_name,
            "customer_email": body.customer_email,
            "subtotal_amount": subtotal_amount,
            "shipping_amount": shipping_amount,
            "total_amount": total_amount,
        }
    except HTTPException:
        raise
    except (PaymentConnectionError, PaymentTimeout) as exc:
        log.error("Payment provider connection failed (%s); check backend network/proxy settings", type(exc).__name__)
        raise HTTPException(503, "Unable to connect to the payment provider. Please try again shortly. No payment has been taken.")
    except (razorpay.errors.BadRequestError, razorpay.errors.ServerError, razorpay.errors.GatewayError):
        log.error("Razorpay rejected order creation; verify matching key pair, account and mode")
        raise HTTPException(503, "Payment provider is unavailable. Please contact us or try again later.")
    except Exception as e:
        log.error("Order creation failed (%s)", type(e).__name__)
        raise HTTPException(status_code=500, detail="Failed to create order")


@api.post("/orders/verify-payment")
async def verify_payment(body: PaymentVerifyRequest):
    """Verify Razorpay payment and mark order as paid."""
    require_payment_configuration()
    try:
        log.info("Payment verification request received", extra={
            "razorpay_order_id": body.razorpay_order_id,
            "razorpay_payment_id": body.razorpay_payment_id,
        })

        # Verify signature
        if not verify_razorpay_signature(body.razorpay_order_id, body.razorpay_payment_id, body.razorpay_signature):
            log.warning("Invalid payment signature for razorpay_order_id=%s", body.razorpay_order_id)
            raise HTTPException(status_code=400, detail="Invalid payment signature")
        
        # Find order by razorpay_order_id
        order = await db.orders.find_one({"razorpay_order_id": body.razorpay_order_id})
        if not order:
            log.warning(f"Order not found: {body.razorpay_order_id}")
            raise HTTPException(status_code=404, detail="Order not found")

        if order.get("payment_status") == "paid":
            if order.get("razorpay_payment_id") != body.razorpay_payment_id:
                raise HTTPException(status_code=409, detail="Order is already paid")
            return {"ok": True, "already_verified": True, "order_id": order.get("order_id"), "total_amount": order.get("total_amount"), "message": "Payment already confirmed"}
        
        # Verify payment with Razorpay API
        try:
            payment = await run_in_threadpool(razorpay_client.payment.fetch, body.razorpay_payment_id, timeout=20)
            log.info("Razorpay payment fetched", extra={
                "razorpay_payment_id": body.razorpay_payment_id,
                "amount": payment.get("amount"),
                "status": payment.get("status"),
            })
            if payment.get("order_id") != body.razorpay_order_id or payment.get("amount") != order.get("total_amount") or payment.get("currency") != "INR":
                log.warning("Payment %s does not match order %s", body.razorpay_payment_id, body.razorpay_order_id)
                raise HTTPException(status_code=400, detail="Payment does not match this order")
            if payment.get("status") != "captured":
                log.warning("Payment not captured for razorpay_payment_id=%s", body.razorpay_payment_id)
                raise HTTPException(status_code=400, detail="Payment not captured")
        except HTTPException:
            raise
        except Exception as e:
            log.error("Razorpay verification failed (%s)", type(e).__name__)
            raise HTTPException(status_code=500, detail="Payment verification failed")
        
        # Atomic transition so concurrent/repeated verifications confirm and notify only once
        marked = await db.orders.update_one(
            {"razorpay_order_id": body.razorpay_order_id, "payment_status": {"$ne": "paid"}},
            {
                "$set": {
                    "payment_status": "paid",
                    "razorpay_payment_id": body.razorpay_payment_id,
                    "razorpay_signature": body.razorpay_signature,
                    "updated_at": datetime.now(timezone.utc).isoformat(),
                }
            }
        )
        if marked.modified_count == 0:
            return {"ok": True, "already_verified": True, "order_id": order.get("order_id"), "total_amount": order.get("total_amount"), "message": "Payment already confirmed"}
        
        # Send Telegram notification (non-blocking)
        order_data = {
            "order_id": order.get("order_id"),
            "razorpay_order_id": order.get("razorpay_order_id"),
            "razorpay_payment_id": body.razorpay_payment_id,
            "payment_method": "Razorpay",
            "payment_status": "paid",
            "customer_name": order.get("customer_name"),
            "customer_email": order.get("customer_email"),
            "customer_phone": order.get("customer_phone"),
            "shipping_address": order.get("shipping_address"),
            "shipping_city": order.get("shipping_city"),
            "shipping_state": order.get("shipping_state"),
            "shipping_pin": order.get("shipping_pin"),
            "items": order.get("items", []),
            "subtotal_amount": order.get("subtotal_amount"),
            "shipping_amount": order.get("shipping_amount"),
            "total_amount": order.get("total_amount"),
            "created_at": order.get("created_at"),
        }
        
        # Fire and forget - don't wait for notifications
        import asyncio
        asyncio.create_task(send_telegram_notification(order_data))
        asyncio.create_task(send_admin_order_email_async(order_data))
        
        log.info("Payment verification completed and notifications queued", extra={
            "order_id": order.get("order_id"),
            "razorpay_order_id": order.get("razorpay_order_id"),
            "razorpay_payment_id": body.razorpay_payment_id,
        })
        
        return {
            "ok": True,
            "order_id": order.get("order_id"),
            "message": "Payment verified and order confirmed",
            "total_amount": order.get("total_amount"),
        }
    except HTTPException:
        raise
    except Exception as e:
        log.error("Payment verification failed (%s)", type(e).__name__)
        raise HTTPException(status_code=500, detail="Payment verification failed")


@api.get("/admin/orders")
async def list_orders(admin: dict = Depends(get_current_admin)):
    """Admin: List all orders."""
    items = await db.orders.find({}, {"_id": 0}).sort([("created_at", -1)]).limit(500).to_list(500)
    return {"items": items}


@api.get("/admin/orders/{order_id}")
async def get_order_detail(order_id: str, admin: dict = Depends(get_current_admin)):
    """Admin: Get order details."""
    order = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    return order


app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("jlt")


@app.on_event("startup")
async def startup():
    if payment_configuration_error():
        log.warning("Online payments disabled: %s", payment_configuration_error())
    await db.products.create_index("slug", unique=True)
    await db.products.create_index("brand_inspiration")
    await db.products.create_index("scent_family")
    await db.products.create_index("moods")
    await db.users.create_index("email", unique=True)
    await db.orders.create_index("order_id", unique=True)
    await db.orders.create_index("razorpay_order_id", unique=True)
    await db.orders.create_index("customer_email")
    await db.orders.create_index("payment_status")
    try:
        await db.orders.drop_index("created_at_1")
    except Exception:
        pass

    admin_email = os.environ["ADMIN_EMAIL"].lower()
    admin_password = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "email": admin_email, "password_hash": hash_password(admin_password),
            "name": "JLT Admin", "role": "admin",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        log.info("Seeded admin user: %s", admin_email)
    if os.environ.get("APP_ENV", "development") != "production" and os.environ.get("ENABLE_DEV_ADMIN") == "true":
        dev_username = os.environ["DEV_ADMIN_USERNAME"].strip().lower()
        dev_password = os.environ["DEV_ADMIN_PASSWORD"]
        dev_email = os.environ.get("DEV_ADMIN_EMAIL", "admin@local.jlt")
        dev_user = await db.users.find_one({"email": dev_email})
        updates = {"username":dev_username, "name":"JLT Development Admin", "role":"admin"}
        if not dev_user or not verify_password(dev_password,dev_user["password_hash"]):
            updates["password_hash"] = hash_password(dev_password)
        await db.users.update_one({"email":dev_email},{"$set":updates},upsert=True)

@app.on_event("shutdown")
async def shutdown():
    client.close()
